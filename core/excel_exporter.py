"""Excel出力モジュール."""

import logging
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)

HEADER_FILL = PatternFill(start_color="D5E8F0", end_color="D5E8F0", fill_type="solid")
HEADER_FONT = Font(bold=True)
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)


class ExcelExporter:
    """Excel形式でのデータ出力を行うクラス."""

    def create(
        self,
        filepath: str | Path,
        data: list[dict],
        headers: list[str] | None = None,
        options: dict | None = None,
    ) -> Path:
        """新規Excelファイルを作成してデータを書き込む.

        Args:
            filepath: 出力ファイルパス
            data: データの辞書リスト
            headers: ヘッダーリスト（省略時はdataのキーから生成）
            options: 書式設定オプション

        Returns:
            出力されたファイルのPath
        """
        filepath = Path(filepath)
        options = options or {}
        wb = Workbook()
        ws = wb.active
        ws.title = "データ"

        if not data:
            wb.save(str(filepath))
            return filepath

        if headers is None:
            headers = list(data[0].keys())

        # ヘッダー行書き込み
        for col_idx, header in enumerate(headers, 1):
            ws.cell(row=1, column=col_idx, value=header)

        # データ行書き込み
        for row_idx, row_data in enumerate(data, 2):
            for col_idx, header in enumerate(headers, 1):
                value = row_data.get(header, "")
                ws.cell(row=row_idx, column=col_idx, value=value)

        # 書式設定
        if options.get("format_header", True):
            self.format_header(ws, len(headers))

        if options.get("auto_column_width", True):
            self.auto_column_width(ws)

        filepath.parent.mkdir(parents=True, exist_ok=True)

        try:
            wb.save(str(filepath))
        except PermissionError:
            alt_path = filepath.with_name(
                f"{filepath.stem}_{datetime.now().strftime('%H%M%S')}{filepath.suffix}"
            )
            logger.warning(
                "ファイルがロックされています。代替パスに保存: %s", alt_path
            )
            wb.save(str(alt_path))
            return alt_path

        logger.info("Excel出力完了: %s (%d件)", filepath, len(data))
        return filepath

    def append(
        self,
        filepath: str | Path,
        data: list[dict],
        headers: list[str] | None = None,
    ) -> Path:
        """既存Excelファイルにデータを追記する.

        ファイルが存在しない場合は新規作成する.
        """
        filepath = Path(filepath)
        if not filepath.exists():
            return self.create(filepath, data, headers)

        wb = load_workbook(str(filepath))
        ws = wb.active

        if not data:
            wb.save(str(filepath))
            return filepath

        if headers is None:
            headers = list(data[0].keys())

        start_row = ws.max_row + 1

        for row_idx, row_data in enumerate(data, start_row):
            for col_idx, header in enumerate(headers, 1):
                value = row_data.get(header, "")
                ws.cell(row=row_idx, column=col_idx, value=value)

        try:
            wb.save(str(filepath))
        except PermissionError:
            alt_path = filepath.with_name(
                f"{filepath.stem}_{datetime.now().strftime('%H%M%S')}{filepath.suffix}"
            )
            logger.warning(
                "ファイルがロックされています。代替パスに保存: %s", alt_path
            )
            wb.save(str(alt_path))
            return alt_path

        logger.info("Excel追記完了: %s (%d件追加)", filepath, len(data))
        return filepath

    def format_header(self, ws, num_columns: int):
        """ヘッダー行に書式を適用する."""
        for col_idx in range(1, num_columns + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.border = THIN_BORDER
            cell.alignment = Alignment(horizontal="center")

    def auto_column_width(self, ws):
        """列幅を内容に合わせて自動調整する."""
        for column_cells in ws.columns:
            max_length = 0
            column_letter = get_column_letter(column_cells[0].column)
            for cell in column_cells:
                try:
                    value = str(cell.value) if cell.value is not None else ""
                    # 日本語文字は幅2として計算
                    length = sum(2 if ord(c) > 127 else 1 for c in value)
                    max_length = max(max_length, length)
                except Exception:
                    pass
            adjusted_width = min(max_length + 4, 60)
            ws.column_dimensions[column_letter].width = adjusted_width

    def add_log_sheet(self, wb: Workbook, log_entries: list[dict]):
        """ログシートを追加する."""
        ws = wb.create_sheet(title="実行ログ")
        log_headers = ["日時", "レベル", "メッセージ"]

        for col_idx, header in enumerate(log_headers, 1):
            ws.cell(row=1, column=col_idx, value=header)

        self.format_header(ws, len(log_headers))

        for row_idx, entry in enumerate(log_entries, 2):
            ws.cell(row=row_idx, column=1, value=entry.get("timestamp", ""))
            ws.cell(row=row_idx, column=2, value=entry.get("level", ""))
            ws.cell(row=row_idx, column=3, value=entry.get("message", ""))

        self.auto_column_width(ws)
