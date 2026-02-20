"""ExcelExporter モジュールのテスト."""

import os
import tempfile
import unittest
from pathlib import Path

from core.excel_exporter import ExcelExporter


class TestExcelExporter(unittest.TestCase):
    """ExcelExporter クラスのテスト."""

    def setUp(self):
        self.exporter = ExcelExporter()
        self.temp_dir = tempfile.mkdtemp()

    def test_create_basic(self):
        """基本的なExcel作成テスト."""
        data = [
            {"タイトル": "記事1", "URL": "https://example.com/1"},
            {"タイトル": "記事2", "URL": "https://example.com/2"},
        ]
        filepath = Path(self.temp_dir) / "test_output.xlsx"
        result = self.exporter.create(filepath, data)
        self.assertTrue(result.exists())

    def test_create_empty(self):
        """空データの場合のテスト."""
        filepath = Path(self.temp_dir) / "empty.xlsx"
        result = self.exporter.create(filepath, [])
        self.assertTrue(result.exists())

    def test_append(self):
        """追記モードのテスト."""
        data1 = [{"名前": "A"}]
        data2 = [{"名前": "B"}, {"名前": "C"}]
        filepath = Path(self.temp_dir) / "append_test.xlsx"

        self.exporter.create(filepath, data1, headers=["名前"])
        self.exporter.append(filepath, data2, headers=["名前"])

        from openpyxl import load_workbook
        wb = load_workbook(str(filepath))
        ws = wb.active
        # ヘッダー(1行) + data1(1行) + data2(2行) = 4行
        self.assertEqual(ws.max_row, 4)

    def test_create_with_format(self):
        """書式付き作成のテスト."""
        data = [{"A": "1", "B": "2"}]
        filepath = Path(self.temp_dir) / "formatted.xlsx"
        options = {"format_header": True, "auto_column_width": True}
        result = self.exporter.create(filepath, data, options=options)
        self.assertTrue(result.exists())

    def tearDown(self):
        """テスト後のクリーンアップ."""
        import shutil
        shutil.rmtree(self.temp_dir, ignore_errors=True)


if __name__ == "__main__":
    unittest.main()
